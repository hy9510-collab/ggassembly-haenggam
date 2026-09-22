# -*- coding: utf-8 -*-
"""
행감 질의·자료요구 → 한글(아래아한글)·워드에서 열리는 .docx 문서 생성
 - 입력: data/qa_data.js, data/units_data.js
 - 출력: 문서/2026_행감_질의서_전체.docx
         문서/2026_행감_자료요구목록.docx   (2025 요구자료 목록과 같은 표 형식)
         문서/기관별/질의서_{기관}.docx
 - 실행: 질의서_문서만들기.bat 더블클릭
"""
import json, os, sys
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, '문서')
NAVY = RGBColor(0x0B, 0x3A, 0x63)
BLUE = RGBColor(0x18, 0x5F, 0xA5)
GRAY = RGBColor(0x6B, 0x72, 0x80)
RED = RGBColor(0xB9, 0x1C, 0x1C)


def load_js(name, var):
    t = open(os.path.join(BASE, 'data', name), encoding='utf-8').read()
    return json.loads(t.split('\nwindow.' + var + ' = ', 1)[1].rstrip().rstrip(';'))


def font(run, size=11, bold=False, color=None, name='맑은 고딕'):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), name)
    if color is not None:
        run.font.color.rgb = color


def para(doc, text='', size=11, bold=False, color=None, align=None, after=4, before=0, indent=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after, pf.space_before, pf.line_spacing = Pt(after), Pt(before), 1.35
    if indent is not None:
        pf.left_indent = Cm(indent)
    if text:
        font(p.add_run(text), size, bold, color)
    return p


def shade(cell, hexcolor):
    tcPr = cell._element.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcolor)
    tcPr.append(sh)


def cell_text(cell, lines, size=9.5, bold_first=False, align=None):
    cell.text = ''
    for i, ln in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        if align:
            p.alignment = align
        font(p.add_run(ln), size, bold_first and i == 0)


def new_doc():
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin = s.bottom_margin = Cm(1.8)
    st = doc.styles['Normal']
    st.font.name = '맑은 고딕'
    st.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    return doc


def cover(doc, title, sub):
    para(doc, '경기도의회 문화체육관광위원회', 12, True, BLUE, WD_ALIGN_PARAGRAPH.CENTER, after=6)
    para(doc, title, 20, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER, after=6)
    para(doc, sub, 10, False, GRAY, WD_ALIGN_PARAGRAPH.CENTER, after=16)


def write_unit(doc, uid, U, QA, fmap):
    u, a = U[uid], QA['units'][uid]
    para(doc, u['name'], 15, True, NAVY, before=10, after=4)
    para(doc, '□ 쟁점 요약', 11.5, True, BLUE, after=2)
    para(doc, a['summary'], 10.5, indent=0.4, after=6)
    para(doc, '□ 질의', 11.5, True, BLUE, after=2)
    for q in a['questions']:
        tag = '【중점】 ' if q['pri'] == '상' else ''
        p = para(doc, '', after=1, indent=0.2)
        font(p.add_run(f"{q['no']}. {tag}"), 10.5, True, RED if q['pri'] == '상' else None)
        font(p.add_run(q['topic']), 10.5, True)
        para(doc, q['q'], 10.5, indent=0.8, after=1)
        refs = []
        if q.get('bg'):
            refs.append('배경: ' + q['bg'])
        fs = [f"{k}({fmap.get((uid, k), '')})" for k in q.get('findings', [])]
        if fs:
            refs.append('2025 지적: ' + ', '.join(fs))
        if q.get('reqs'):
            refs.append('관련 자료요구: ' + ', '.join(q['reqs']))
        for r in refs:
            para(doc, '· ' + r, 9.5, color=GRAY, indent=0.8, after=0)
        para(doc, '', after=3)
    if a['requests']:
        para(doc, '□ 자료요구', 11.5, True, BLUE, after=2)
        for r in a['requests']:
            para(doc, f"[{r['no']}] ∘ {r['title']}", 10.5, True, indent=0.2, after=0)
            for it in r['items']:
                para(doc, '- ' + it, 10, indent=0.9, after=0)
            para(doc, '', after=2)


def req_table(doc, rows, member='상임위원'):
    t = doc.add_table(rows=1, cols=4)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(1.3), Cm(2.0), Cm(11.2), Cm(2.3)]
    for i, h in enumerate(['연번', '의원명', '자     료     명', '비  고']):
        c = t.rows[0].cells[i]
        cell_text(c, [h], 10, True, WD_ALIGN_PARAGRAPH.CENTER)
        shade(c, 'DCE6F1')
        c.width = widths[i]
    for n, r in enumerate(rows, 1):
        cells = t.add_row().cells
        cell_text(cells[0], [str(n)], 9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(cells[1], [member], 9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(cells[2], ['∘ ' + r['title']] + [' - ' + x for x in r['items']], 9.5, bold_first=True)
        cell_text(cells[3], [r['no']], 8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        for i, w in enumerate(widths):
            cells[i].width = w
    para(doc, '', after=6)


def main():
    QA = load_js('qa_data.js', 'QA')
    UN = load_js('units_data.js', 'UNITS')
    U = {u['id']: u for u in UN['units']}
    fmap = {(u['id'], f"{f['type']}-{f['no']}"): ('완료' if f['status'] == '완료' else '추진 중') for u in UN['units'] for f in u['findings']}
    order = [u['id'] for u in UN['units']]
    today = datetime.now().strftime('%Y. %m. %d.')
    os.makedirs(os.path.join(OUT, '기관별'), exist_ok=True)

    # 1) 전체 질의서
    doc = new_doc()
    cover(doc, '2026년 행정사무감사 질의서(안)', f'문화체육관광국 및 산하기관 · 작성 {today} · 자료 기준 {QA["updated"]}')
    para(doc, 'Ⅰ. 행감 추천 의제', 14, True, NAVY, after=4)
    for a in QA.get('agenda', []):
        p = para(doc, '', after=1)
        font(p.add_run(f"{a['rank']}. "), 11, True, RED)
        font(p.add_run(a['title']), 11, True)
        para(doc, a['why'], 10, indent=0.6, after=0)
        para(doc, '· 관련: ' + ', '.join(U[x]['name'] for x in a['units']) + ' / 진행: ' + a['how'], 9.5, color=GRAY, indent=0.6, after=4)
    if QA.get('bench'):
        doc.add_page_break()
        para(doc, 'Ⅱ. 정부·타 시도 관련 자료', 14, True, NAVY, after=4)
        for b in QA['bench']:
            para(doc, f"[{b['id']}] {b['topic']}", 11, True, BLUE, before=4, after=1)
            for g in b['gov']:
                para(doc, '· (정부) ' + g, 10, indent=0.5, after=0)
            for o in b['other']:
                para(doc, f"· ({o['r']}) {o['t']}", 10, indent=0.5, after=0)
            para(doc, '· (경기도) ' + b['gg'], 10, indent=0.5, after=0)
            para(doc, '▶ 시사점: ' + b['impl'], 10, True, indent=0.5, after=0)
            para(doc, '출처: ' + ' / '.join(f"{s['src']} {s['url']}" for s in b['src']), 8.5, color=GRAY, indent=0.5, after=3)
    doc.add_page_break()
    para(doc, 'Ⅲ. 기관·부서별 질의 및 자료요구', 14, True, NAVY, after=4)
    for i, uid in enumerate(order):
        if i:
            doc.add_page_break()
        write_unit(doc, uid, U, QA, fmap)
    f1 = os.path.join(OUT, '2026_행감_질의서_전체.docx')
    doc.save(f1)

    # 2) 자료요구 목록(2025 표 형식)
    doc = new_doc()
    cover(doc, '2026년 행정사무감사 요구자료 목록(안)', f'문화체육관광위원회 · 작성 {today}')
    para(doc, '【 공 통 】', 12, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER, after=4)
    req_table(doc, QA['common'])
    for uid in order:
        rs = QA['units'][uid]['requests']
        if not rs:
            continue
        para(doc, f"【 {U[uid]['name']} 】", 12, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER, before=6, after=4)
        req_table(doc, rs)
    f2 = os.path.join(OUT, '2026_행감_자료요구목록.docx')
    doc.save(f2)

    # 3) 기관별 질의서
    n = 0
    for uid in order:
        doc = new_doc()
        cover(doc, f"{U[uid]['name']} 질의서(안)", f'2026년 행정사무감사 · 작성 {today}')
        write_unit(doc, uid, U, QA, fmap)
        name = U[uid]['name'].replace(' (총괄)', '_총괄').replace('/', '_')
        doc.save(os.path.join(OUT, '기관별', f'질의서_{name}.docx'))
        n += 1
    print('[완료]', os.path.relpath(f1, BASE))
    print('[완료]', os.path.relpath(f2, BASE))
    print(f'[완료] 문서/기관별/ 질의서 {n}개')


if __name__ == '__main__':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass
    main()
